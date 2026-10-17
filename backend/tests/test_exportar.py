import copy
import io
import json
import re
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

from exportar.docx import generar_docx
from exportar.pdf import generar_pdf

FIXTURE_DEMO = Path(__file__).parent / "fixtures" / "eventos_demo.json"


def reporte_demo() -> dict:
    eventos = json.loads(FIXTURE_DEMO.read_text(encoding="utf-8"))
    return copy.deepcopy(eventos[-1]["resultado"]["reporte"])


def reporte_completo() -> dict:
    """El de la demo, más una prueba respaldada por una guía, para recorrer todas las ramas."""
    r = reporte_demo()
    r["pruebas"].append({
        "nombre": "PHQ-9", "motivo": "Tamizaje de depresión",
        "guia": {"titulo": "Guía de depresión", "url": "https://www.imss.gob.mx/guias/depresion"},
        "cita": "se recomienda el uso de escalas validadas",
    })
    return r


def texto_docx(datos: bytes) -> str:
    d = Document(io.BytesIO(datos))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        partes += [c.text for f in t.rows for c in f.cells]
    return "\n".join(partes)


def texto_pdf(datos: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(datos)).pages)


# ----------------------------------------------------------------------------- Word
def test_docx_usa_arial_12_e_interlineado_1_5():
    d = Document(io.BytesIO(generar_docx(reporte_completo())))
    normal = d.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5
    for estilo in ("Heading 1", "Heading 2"):
        assert d.styles[estilo].font.name == "Arial"
        assert str(d.styles[estilo].font.color.rgb) == "2F4B7C"


def test_docx_contiene_secciones_y_aviso():
    r = reporte_completo()
    texto = texto_docx(generar_docx(r))
    for esperado in (
        "Ψ-DIAG con IA", "Datos del caso", "Laura Mendoza", "Síntomas por categoría", "Tristeza", "Hipótesis diagnósticas",
        "6A70", "Trastorno depresivo de episodio único", "Ánimo deprimido casi todos los días", "Ambigüedades", "Diferenciales",
        "Alternativas consultadas y descartadas", "6A71", "Datos faltantes", "Pruebas complementarias sugeridas",
        "Inventario de Depresión de Beck", "sin referencia en las guías consultadas", "PHQ-9", "se recomienda el uso de escalas validadas",
        "Guías clínicas consultadas", "IMSS (México) · Guías de práctica clínica", r["aviso"],
    ):
        assert esperado in texto, esperado


def test_docx_enlaza_a_la_oms_y_a_las_guias():
    d = Document(io.BytesIO(generar_docx(reporte_completo())))
    ligas = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://icd.who.int/browse/2024-01/mms/es#1001" in ligas
    assert "https://www.imss.gob.mx/guias/depresion" in ligas


def test_docx_sin_guias_no_dice_sin_referencia():
    r = reporte_demo()
    r["guias"] = []
    assert "sin referencia" not in texto_docx(generar_docx(r))


# ----------------------------------------------------------------------------- PDF
def test_pdf_empieza_con_pdf_y_tiene_aviso_y_numero_de_pagina():
    r = reporte_completo()
    datos = generar_pdf(r)
    assert datos.startswith(b"%PDF")
    texto = texto_pdf(datos)
    assert " ".join(r["aviso"].split()) in " ".join(texto.split())
    assert "Herramienta de apoyo; no sustituye el criterio clínico" in texto
    assert re.search(r"Página 1 de \d+", texto)
    assert "DIAG con" in texto and "IA" in texto
    assert "Hipótesis diagnósticas" in texto and "6A70" in texto


def test_pdf_incrusta_nunito():
    lector = PdfReader(io.BytesIO(generar_pdf(reporte_completo())))
    fuentes = {}
    for pagina in lector.pages:
        for f in (pagina["/Resources"].get("/Font") or {}).values():
            f = f.get_object()
            assert "/FontDescriptor" in f, f"fuente sin incrustar: {f['/BaseFont']}"
            fuentes[str(f["/BaseFont"])] = f["/FontDescriptor"].get_object()
    assert fuentes and all("Nunito" in nombre for nombre in fuentes)
    assert all("/FontFile2" in d for d in fuentes.values())


def test_pdf_tiene_ligas_clicables():
    lector = PdfReader(io.BytesIO(generar_pdf(reporte_completo())))
    anotaciones = [a.get_object() for p in lector.pages for a in p.get("/Annots", [])]
    destinos = {a["/A"]["/URI"] for a in anotaciones if "/A" in a}   # las del índice son internas (/Dest), sin URI
    assert "https://icd.who.int/browse/2024-01/mms/es#1001" in destinos
    assert "https://www.imss.gob.mx/guias/depresion" in destinos


def test_pdf_largo_numera_varias_paginas():
    r = reporte_demo()
    r["datos_faltantes"] = [f"Dato faltante número {i} con texto de relleno suficiente para ocupar espacio" for i in range(150)]
    texto = texto_pdf(generar_pdf(r))
    assert "Página 1 de " in texto and "Página 2 de " in texto


# ----------------------------------------------------------------------------- texto hostil
def reporte_hostil() -> dict:
    hostil = "<script>alert(1)</script> &amp; ñandú áéíóú 😀🔥 \x00\x07 {{x}} <b>"
    larga = "a" * 500
    return {
        "fecha": None,
        "caso": {"paciente_ref": hostil, "edad": None, "sexo": None, "temporalidad": larga},
        "sintomas": {"cognitivos": [hostil, larga], "emocionales": [], "conductuales": None, "afectacion_social": ["😀"]},
        "hipotesis": [
            {"codigo": "6A70", "titulo": hostil, "url": "javascript:alert(1)", "justificacion": larga,
             "criterios_cumplidos": [], "criterios_no_cumplidos": [hostil, None]},
            {"codigo": None, "titulo": None, "url": None, "justificacion": None, "criterios_cumplidos": None, "criterios_no_cumplidos": []},
        ],
        "ambiguedades": [], "diferenciales": [hostil],
        "alternativas": [{"codigo": "X", "titulo": hostil, "url": "https://icd.who.int/\"><script>"}],
        "datos_faltantes": [larga, ""],
        "pruebas": [
            {"nombre": hostil, "motivo": larga, "guia": None, "cita": ""},
            {"nombre": None, "motivo": None, "guia": None, "cita": None},
            {"nombre": "p", "motivo": "m", "guia": {"titulo": hostil, "url": "https://x.org/?a=1&b=<2>"}, "cita": hostil},
        ],
        "guias": [{"tipo": "oms", "etiqueta": hostil, "titulo": larga, "url": "https://x.org", "cita": None}],
        "aviso": hostil,
    }


def reporte_vacio() -> dict:
    return {"fecha": "no es fecha", "caso": {}, "sintomas": {}, "hipotesis": [], "ambiguedades": [], "diferenciales": [],
            "alternativas": [], "datos_faltantes": [], "pruebas": [], "guias": [], "aviso": ""}


@pytest.mark.parametrize("formato", ["docx", "pdf"])
@pytest.mark.parametrize("fabrica", [reporte_hostil, reporte_vacio])
def test_exportadores_aceptan_texto_hostil(formato, fabrica):
    if formato == "docx":
        datos = generar_docx(fabrica())
        Document(io.BytesIO(datos))
    else:
        datos = generar_pdf(fabrica())
        assert len(PdfReader(io.BytesIO(datos)).pages) >= 1
    assert datos


def test_docx_conserva_el_texto_hostil_como_texto_y_sin_ligas_peligrosas():
    d = Document(io.BytesIO(generar_docx(reporte_hostil())))
    assert "<script>alert(1)</script>" in texto_docx(generar_docx(reporte_hostil()))
    ligas = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert not any(l.startswith("javascript:") for l in ligas)


def test_pdf_no_crea_ligas_peligrosas_y_muestra_el_texto_hostil():
    datos = generar_pdf(reporte_hostil())
    lector = PdfReader(io.BytesIO(datos))
    destinos = [str(a.get_object()["/A"]["/URI"]) for p in lector.pages for a in p.get("/Annots", []) if "/A" in a.get_object()]
    assert not any(d.startswith("javascript:") for d in destinos)
    assert "<script>alert(1)</script>" in texto_pdf(datos)


# ----------------------------------------------------------------------------- API
import logging  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402

import servidor as modulo_app  # noqa: E402

TIPOS = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}


@pytest.fixture(autouse=True)
def entorno_api(monkeypatch):
    for variable in ("ACCESS_CODE", "VERCEL"):
        monkeypatch.delenv(variable, raising=False)
    modulo_app._visitas.clear()
    modulo_app._fallos.clear()


@pytest.fixture
def cliente():
    return TestClient(modulo_app.app)


@pytest.mark.parametrize("formato", ["docx", "pdf"])
def test_api_devuelve_el_archivo_con_sus_cabeceras(cliente, formato):
    r = cliente.post(f"/api/exportar/{formato}", json=reporte_demo())
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(TIPOS[formato])
    assert r.headers["content-disposition"] == f'attachment; filename="reporte-2026-10-15.{formato}"'
    assert r.headers["cache-control"] == "no-store"
    assert r.headers["x-content-type-options"] == "nosniff"
    if formato == "pdf":
        assert r.content.startswith(b"%PDF")
    else:
        Document(io.BytesIO(r.content))


def test_api_fecha_invalida_usa_la_de_hoy_y_no_filtra_texto_al_nombre(cliente):
    from datetime import date
    r = cliente.post("/api/exportar/pdf", json={**reporte_demo(), "fecha": '"; evil=1\r\n'})
    assert r.headers["content-disposition"] == f'attachment; filename="reporte-{date.today().isoformat()}.pdf"'


@pytest.mark.parametrize("formato", ["md", "html", "DOCX", "docx.exe"])
def test_api_otro_formato_responde_404_con_la_forma_de_error(cliente, formato):
    r = cliente.post(f"/api/exportar/{formato}", json=reporte_demo())
    assert r.status_code == 404
    cuerpo = r.json()
    assert cuerpo["ok"] is False and cuerpo["error"]["codigo"] == "validacion"


def test_api_exige_el_codigo_de_acceso(cliente, monkeypatch):
    monkeypatch.setenv("ACCESS_CODE", "secreto")
    assert cliente.post("/api/exportar/pdf", json=reporte_demo()).status_code == 401
    assert cliente.post("/api/exportar/pdf", json=reporte_demo(), headers={"x-access-code": "mal"}).status_code == 401
    assert cliente.post("/api/exportar/pdf", json=reporte_demo(), headers={"x-access-code": "secreto"}).status_code == 200


def test_api_intentos_con_codigo_incorrecto_terminan_en_429(cliente, monkeypatch):
    monkeypatch.setenv("ACCESS_CODE", "secreto")
    codigos = [cliente.post("/api/exportar/pdf", json=reporte_demo(), headers={"x-access-code": "mal"}).status_code for _ in range(7)]
    assert codigos[:5] == [401] * 5 and codigos[-1] == 429


def test_api_limita_solicitudes_por_minuto(cliente, monkeypatch):
    monkeypatch.setattr(modulo_app, "MAX_POR_MINUTO", 2)
    codigos = [cliente.post("/api/exportar/docx", json=reporte_demo()).status_code for _ in range(3)]
    assert codigos == [200, 200, 429]


def test_api_rechaza_cuerpo_mayor_a_200_kb(cliente):
    grande = {**reporte_demo(), "relleno": "x" * 200_001}
    r = cliente.post("/api/exportar/pdf", json=grande)
    assert r.status_code == 413 and r.json()["error"]["codigo"] == "validacion"


@pytest.mark.parametrize("cuerpo", [b"no es json", b"[1, 2]", b'"texto"', b"{}", b'{"caso": {}}', b""])
def test_api_rechaza_cuerpos_que_no_son_un_reporte(cliente, cuerpo):
    r = cliente.post("/api/exportar/docx", content=cuerpo, headers={"content-type": "application/json"})
    assert r.status_code == 400 and r.json()["error"]["codigo"] == "validacion"


def test_api_acepta_texto_hostil(cliente):
    assert cliente.post("/api/exportar/pdf", json=reporte_hostil()).status_code == 200
    assert cliente.post("/api/exportar/docx", json=reporte_hostil()).status_code == 200


def test_api_no_registra_el_contenido_del_reporte(cliente, caplog):
    r = reporte_demo()
    r["caso"]["paciente_ref"] = "NOMBRE-SECRETO-XYZ"
    with caplog.at_level(logging.INFO, logger="psico"):
        modulo_app.log.propagate = True
        try:
            assert cliente.post("/api/exportar/pdf", json=r).status_code == 200
        finally:
            modulo_app.log.propagate = False
    salida = " ".join(caplog.messages)
    assert "NOMBRE-SECRETO-XYZ" not in salida and "Laura" not in salida
    linea = json.loads(caplog.messages[-1])
    assert linea["formato"] == "pdf" and set(linea) == {"id", "formato", "ms", "estado"}


# ----------------------------------------------------------------------------- sustitutos UTF-16 sueltos
def reporte_con_sustitutos() -> dict:
    s = "\ud800"
    r = reporte_completo()
    r["fecha"] = f"2026-10-15{s}"
    r["caso"]["paciente_ref"] = f"Ana{s}"
    r["sintomas"]["emocionales"] = [f"{s}tristeza", s]
    r["hipotesis"][0].update(titulo=f"T{s}", codigo=f"6A{s}", justificacion=f"j{s}\udfff",
                             criterios_cumplidos=[f"c{s}"], criterios_no_cumplidos=[s + "x"])
    r["hipotesis"][0]["url"] = f"https://icd.who.int/{s}"
    r["ambiguedades"], r["diferenciales"], r["datos_faltantes"] = [f"a{s}"], [f"d{s}"], [f"f{s}"]
    r["pruebas"][0].update(nombre=f"p{s}", motivo=f"m{s}", cita=f"c{s}")
    r["pruebas"][-1]["guia"]["titulo"] = f"g{s}"
    r["guias"][0].update(etiqueta=f"e{s}", titulo=f"t{s}", cita=f"c{s}")
    r["aviso"] = f"aviso{s}"
    return r


@pytest.mark.parametrize("formato", ["docx", "pdf"])
def test_exportadores_aceptan_sustitutos_utf16_sueltos(formato):
    if formato == "docx":
        Document(io.BytesIO(generar_docx(reporte_con_sustitutos())))
    else:
        assert PdfReader(io.BytesIO(generar_pdf(reporte_con_sustitutos()))).pages


@pytest.mark.parametrize("formato", ["docx", "pdf"])
def test_api_acepta_sustitutos_utf16_sueltos(cliente, formato):
    cuerpo = json.dumps(reporte_con_sustitutos())   # ensure_ascii deja «\ud800» escapado: json.loads lo acepta
    r = cliente.post(f"/api/exportar/{formato}", content=cuerpo, headers={"content-type": "application/json"})
    assert r.status_code == 200
    assert r.headers["content-disposition"] == f'attachment; filename="reporte-2026-10-15.{formato}"'
    Document(io.BytesIO(r.content)) if formato == "docx" else PdfReader(io.BytesIO(r.content)).pages


def test_docx_lleva_metadatos_propios_y_no_los_de_la_plantilla():
    from datetime import datetime, timedelta, timezone
    propiedades = Document(io.BytesIO(generar_docx(reporte_demo()))).core_properties
    assert propiedades.title == "Reporte de apoyo al psicodiagnóstico"
    assert propiedades.author == propiedades.last_modified_by == "Ψ-DIAG con IA"
    ahora = datetime.now(timezone.utc)
    for marca in (propiedades.created, propiedades.modified):
        assert abs(ahora - marca) < timedelta(minutes=5)


# ----------------------------------------------------------------------------- conclusión e índices
def reporte_con_conclusion() -> dict:
    r = reporte_completo()
    r["conclusion"] = "Resumen del cuadro y siguientes pasos para el profesional."
    return r


def titulos_del_panel(lector) -> list:
    def recorrer(items):
        for x in items:
            if isinstance(x, list):
                yield from recorrer(x)
            else:
                yield x.title, lector.get_destination_page_number(x) + 1
    return list(recorrer(lector.outline))


def test_pdf_tiene_indice_con_ligas_y_marcadores_en_la_pagina_correcta():
    lector = PdfReader(io.BytesIO(generar_pdf(reporte_con_conclusion())))
    texto = lector.pages[0].extract_text()
    assert "Índice" in texto and "1. Datos del caso" in texto and "Conclusión, recomendaciones y siguientes pasos" in texto
    panel = dict(titulos_del_panel(lector))
    assert "1. Datos del caso" in panel and "9. Aviso" in panel
    assert panel["1. Datos del caso"] >= 2 and panel["9. Aviso"] >= panel["1. Datos del caso"]
    assert panel["9. Aviso"] == len(lector.pages)   # el aviso cierra el documento
    internas = [a.get_object() for a in lector.pages[0].get("/Annots", []) if "/Dest" in a.get_object()]
    assert len(internas) >= 9   # una liga por entrada del índice


def test_pdf_largo_apunta_cada_titulo_a_su_pagina():
    r = reporte_con_conclusion()
    r["datos_faltantes"] = [f"Dato faltante número {i} con texto de relleno suficiente para ocupar espacio" for i in range(150)]
    panel = dict(titulos_del_panel(PdfReader(io.BytesIO(generar_pdf(r)))))
    assert panel["6. Datos faltantes por recabar"] < panel["7. Pruebas complementarias sugeridas"] <= panel["9. Aviso"]


def test_pdf_con_titulos_hostiles_no_rompe_el_indice():
    assert generar_pdf(reporte_hostil()).startswith(b"%PDF")


def test_la_conclusion_aparece_en_pdf_y_word_antes_del_aviso():
    r = reporte_con_conclusion()
    assert "Resumen del cuadro y siguientes pasos" in texto_pdf(generar_pdf(r))
    palabra = texto_docx(generar_docx(r))
    assert palabra.index("8. Conclusión, recomendaciones y siguientes pasos") < palabra.index("9. Aviso")
    assert "Resumen del cuadro y siguientes pasos" in palabra


def test_sin_conclusion_el_aviso_sigue_siendo_la_seccion_8():
    r = reporte_completo()
    r.pop("conclusion", None)
    assert "8. Aviso" in texto_docx(generar_docx(r))
    assert "8. Aviso" in texto_pdf(generar_pdf(r))


def test_word_tiene_campo_de_indice_con_ligas_a_los_titulos():
    documento = Document(io.BytesIO(generar_docx(reporte_con_conclusion())))
    xml = documento.element.xml
    assert 'TOC \\o "1-2" \\h \\z \\u' in xml and 'w:dirty="true"' in xml
    assert xml.count("w:fldCharType=\"begin\"") == 1 and xml.count("w:fldCharType=\"end\"") == 1
    anclas = set(re.findall(r'w:hyperlink w:anchor="(_Toc\d+)"', xml))
    marcadores = set(re.findall(r'w:bookmarkStart [^>]*w:name="(_Toc\d+)"', xml))
    assert anclas and anclas == marcadores
    titulos = [p.text for p in documento.paragraphs if p.style.name.startswith("Heading") and p.style.name != "Heading 3"]
    assert "1. Datos del caso" in titulos
    assert "Reporte de apoyo al psicodiagnóstico" not in titulos   # el subtítulo no entra al índice


def test_la_conclusion_en_varios_parrafos_sale_como_parrafos_separados():
    r = reporte_completo()
    r["conclusion"] = "Primero.\n\nSegundo.\n\nTercero."
    palabra = Document(io.BytesIO(generar_docx(r)))
    textos = [p.text for p in palabra.paragraphs]
    assert {"Primero.", "Segundo.", "Tercero."} <= set(textos)
    assert "Primero." in texto_pdf(generar_pdf(r))


def test_pdf_y_word_incluyen_fisiologicos_y_la_seccion_de_enfermedades_y_medicacion():
    r = reporte_completo()
    r["sintomas"] = {**r["sintomas"], "fisiologicos": ["Taquicardia y sudoración"]}
    r["antecedentes_medicos"] = {"condiciones_medicas": ["Diabetes tipo 2"], "medicamentos": ["Metformina"]}
    for texto in (texto_docx(generar_docx(r)), texto_pdf(generar_pdf(r))):
        assert "Fisiológicos" in texto and "Taquicardia y sudoración" in texto
        assert "3. Enfermedades conocidas y medicación" in texto and "Diabetes tipo 2" in texto and "Metformina" in texto


def test_un_reporte_sin_antecedentes_medicos_sigue_exportando():
    r = reporte_completo()
    r.pop("antecedentes_medicos", None)
    assert "3. Enfermedades conocidas y medicación" in texto_docx(generar_docx(r))
    assert generar_pdf(r).startswith(b"%PDF")
