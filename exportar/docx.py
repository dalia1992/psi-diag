"""Reporte en Word: Arial 12 pt, interlineado 1.5, títulos en azul sobrio. El nombre de la app va en texto simple."""
import io

from docx import Document
from docx.enum.text import WD_UNDERLINE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from . import comun
from .comun import dicc, lista, lista_de_dicts, mayuscula_inicial, texto, url_segura

AZUL = RGBColor(0x2F, 0x4B, 0x7C)
FUENTE = "Arial"
NINGUNO = "Ninguno indicado"


def _fuente(estilo, tamano, negrita=False, color=None) -> None:
    estilo.font.name = FUENTE
    estilo.font.size = Pt(tamano)
    estilo.font.bold = negrita
    if color is not None:
        estilo.font.color.rgb = color
    propiedades = estilo.element.get_or_add_rPr()
    fuentes = propiedades.find(qn("w:rFonts"))
    if fuentes is None:
        fuentes = OxmlElement("w:rFonts")
        propiedades.insert(0, fuentes)
    for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fuentes.set(qn(atributo), FUENTE)
    for atributo in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):   # el tema pisaría a Arial
        if fuentes.get(qn(atributo)) is not None:
            del fuentes.attrib[qn(atributo)]


def _estilos(documento) -> None:
    normal = documento.styles["Normal"]
    _fuente(normal, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)
    for nombre, tamano in (("Title", 22), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        _fuente(documento.styles[nombre], tamano, negrita=True, color=AZUL)
        documento.styles[nombre].paragraph_format.line_spacing = 1.5
    _fuente(documento.styles["List Bullet"], 12)
    documento.styles["List Bullet"].paragraph_format.line_spacing = 1.5


def _liga(parrafo, url: str, etiqueta: str) -> None:
    relacion = parrafo.part.relate_to(url, RT.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), relacion)
    corrida = parrafo.add_run(etiqueta)
    corrida.font.color.rgb = AZUL
    corrida.font.underline = WD_UNDERLINE.SINGLE
    enlace.append(corrida._r)
    parrafo._p.append(enlace)


def _vinetas(documento, items, vacio=None) -> None:
    items = lista(items)
    for item in items:
        documento.add_paragraph(mayuscula_inicial(item), style="List Bullet")
    if not items and vacio:
        documento.add_paragraph(vacio, style="List Bullet")


def _nota(documento, contenido: str) -> None:
    documento.add_paragraph().add_run(contenido).italic = True


def _con_liga(documento, etiqueta: str, url: str, resto: str = "", estilo: str | None = None, prefijo: str = "") -> None:
    parrafo = documento.add_paragraph(prefijo, style=estilo)
    liga = url_segura(url)
    if liga:
        _liga(parrafo, liga, etiqueta)
    else:
        parrafo.add_run(etiqueta)
    if resto:
        parrafo.add_run(resto)


def _negrita_y_texto(documento, rotulo: str, contenido: str) -> None:
    parrafo = documento.add_paragraph()
    parrafo.add_run(f"{rotulo} ").bold = True
    parrafo.add_run(contenido)


def generar_docx(reporte: dict) -> bytes:
    reporte = dicc(reporte)
    documento = Document()
    _estilos(documento)
    for seccion in documento.sections:
        seccion.left_margin = seccion.right_margin = Pt(72)

    documento.add_paragraph(comun.NOMBRE_APP, style="Title")
    documento.add_paragraph("Reporte de apoyo al psicodiagnóstico", style="Heading 2")

    documento.add_heading("1. Datos del caso", level=1)
    for rotulo, valor in comun.datos_del_caso(reporte):
        _negrita_y_texto(documento, f"{rotulo}:", valor)

    documento.add_heading("2. Síntomas por categoría", level=1)
    sintomas = dicc(reporte.get("sintomas"))
    for etiqueta, clave in comun.ETIQUETAS_SINTOMAS:
        documento.add_heading(etiqueta, level=2)
        _vinetas(documento, sintomas.get(clave), "Ninguno identificado")

    documento.add_heading("3. Hipótesis diagnósticas (de la más a la menos plausible)", level=1)
    hipotesis = lista_de_dicts(reporte.get("hipotesis"))
    for i, h in enumerate(hipotesis, 1):
        documento.add_heading(f"{i}. {mayuscula_inicial(h.get('titulo')) or 'Sin título'} — {texto(h.get('codigo')) or 'sin código'}", level=2)
        _con_liga(documento, "Abrir en el navegador OMS", h.get("url"))
        _negrita_y_texto(documento, "Justificación:", mayuscula_inicial(h.get("justificacion")) or "No indicada")
        documento.add_paragraph().add_run("Criterios cumplidos:").bold = True
        _vinetas(documento, h.get("criterios_cumplidos"), NINGUNO)
        documento.add_paragraph().add_run("Criterios no cumplidos o sin confirmar:").bold = True
        _vinetas(documento, h.get("criterios_no_cumplidos"), NINGUNO)
    if not hipotesis:
        documento.add_paragraph("Sin hipótesis diagnósticas.")
    if lista(reporte.get("ambiguedades")):
        documento.add_heading("Ambigüedades", level=2)
        _vinetas(documento, reporte.get("ambiguedades"))

    documento.add_heading("4. Diferenciales y exclusiones", level=1)
    _vinetas(documento, reporte.get("diferenciales"), NINGUNO)
    alternativas = lista_de_dicts(reporte.get("alternativas"))
    if alternativas:
        documento.add_heading("Alternativas consultadas y descartadas", level=2)
        for a in alternativas:
            _con_liga(documento, "OMS", a.get("url"), "", "List Bullet",
                      f"{mayuscula_inicial(a.get('titulo')) or 'Sin título'} ({texto(a.get('codigo')) or 'sin código'}) — ")

    documento.add_heading("5. Datos faltantes por recabar", level=1)
    _vinetas(documento, reporte.get("datos_faltantes"), NINGUNO)

    documento.add_heading("6. Pruebas complementarias sugeridas", level=1)
    _nota(documento, "Sugerencias a validar por el profesional. Solo se indica el nombre de la prueba: "
                     "no se incluyen cuestionarios ni formas de calificación.")
    guias = lista_de_dicts(reporte.get("guias"))
    pruebas = lista_de_dicts(reporte.get("pruebas"))
    for p in pruebas:
        parrafo = documento.add_paragraph(style="List Bullet")
        parrafo.add_run(mayuscula_inicial(p.get("nombre")) or "Prueba sin nombre").bold = True
        motivo = mayuscula_inicial(p.get("motivo"))
        if motivo:
            parrafo.add_run(f": {motivo}")
        guia = dicc(p.get("guia"))
        if guia:
            parrafo.add_run(" · Referencia: ")
            titulo, liga = texto(guia.get("titulo")) or "Guía", url_segura(guia.get("url"))
            _liga(parrafo, liga, titulo) if liga else parrafo.add_run(titulo)
            if texto(p.get("cita")):
                parrafo.add_run(f" — «{texto(p.get('cita'))}»")
        elif guias:
            parrafo.add_run(" · (sin referencia en las guías consultadas)").italic = True
    if not pruebas:
        documento.add_paragraph("Ninguna sugerida", style="List Bullet")

    if guias:
        documento.add_heading("Guías clínicas consultadas", level=2)
        _nota(documento, "Referencias para consulta del profesional; el sistema no recomienda tratamientos.")
        etiquetas = list(dict.fromkeys(texto(g.get("etiqueta")) or texto(g.get("tipo")) or "Otras" for g in guias))
        for etiqueta in etiquetas:
            documento.add_paragraph().add_run(etiqueta).bold = True
            for g in guias:
                if (texto(g.get("etiqueta")) or texto(g.get("tipo")) or "Otras") != etiqueta:
                    continue
                cita = texto(g.get("cita"))
                _con_liga(documento, texto(g.get("titulo")) or "Guía", g.get("url"), f" — «{cita}»" if cita else "", "List Bullet")

    documento.add_heading("7. Aviso", level=1)
    documento.add_paragraph(texto(reporte.get("aviso")))
    salida = io.BytesIO()
    documento.save(salida)
    return salida.getvalue()

