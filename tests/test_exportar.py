import copy
import io
import json
import re
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

from exportar.docx import generar_docx
from exportar.pdf import generar_pdf

FIXTURE_DEMO = Path(__file__).parent / "fixtures" / "eventos_demo.json"


def reporte_demo() -> dict:
    eventos = json.loads(FIXTURE_DEMO.read_text(encoding="utf-8"))
    return copy.deepcopy(eventos[-1]["resultado"]["reporte"])


def reporte_completo() -> dict:
    """El de la demo, más una prueba respaldada por una guía, para recorrer todas las ramas."""
    r = reporte_demo()
    r["pruebas"].append({
        "nombre": "PHQ-9", "motivo": "Tamizaje de depresión",
        "guia": {"titulo": "Guía de depresión", "url": "https://www.imss.gob.mx/guias/depresion"},
        "cita": "se recomienda el uso de escalas validadas",
    })
    return r


def texto_docx(datos: bytes) -> str:
    d = Document(io.BytesIO(datos))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        partes += [c.text for f in t.rows for c in f.cells]
    return "\n".join(partes)


def texto_pdf(datos: bytes) -> str:
    return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(datos)).pages)


# ----------------------------------------------------------------------------- Word
def test_docx_usa_arial_12_e_interlineado_1_5():
    d = Document(io.BytesIO(generar_docx(reporte_completo())))
    normal = d.styles["Normal"]
    assert normal.font.name == "Arial"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5
    for estilo in ("Heading 1", "Heading 2"):
        assert d.styles[estilo].font.name == "Arial"
        assert str(d.styles[estilo].font.color.rgb) == "2F4B7C"


def test_docx_contiene_secciones_y_aviso():
    r = reporte_completo()
    texto = texto_docx(generar_docx(r))
    for esperado in (
        "Ψ-DIAG con IA", "Datos del caso", "Laura Mendoza", "Síntomas por categoría", "Tristeza", "Hipótesis diagnósticas",
        "6A70", "Trastorno depresivo de episodio único", "Ánimo deprimido casi todos los días", "Ambigüedades", "Diferenciales",
        "Alternativas consultadas y descartadas", "6A71", "Datos faltantes", "Pruebas complementarias sugeridas",
        "Inventario de Depresión de Beck", "sin referencia en las guías consultadas", "PHQ-9", "se recomienda el uso de escalas validadas",
        "Guías clínicas consultadas", "IMSS (México) · Guías de práctica clínica", r["aviso"],
    ):
        assert esperado in texto, esperado


def test_docx_enlaza_a_la_oms_y_a_las_guias():
    d = Document(io.BytesIO(generar_docx(reporte_completo())))
    ligas = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert "https://icd.who.int/browse/2024-01/mms/es#1001" in ligas
    assert "https://www.imss.gob.mx/guias/depresion" in ligas


def test_docx_sin_guias_no_dice_sin_referencia():
    r = reporte_demo()
    r["guias"] = []
    assert "sin referencia" not in texto_docx(generar_docx(r))


# ----------------------------------------------------------------------------- PDF
def test_pdf_empieza_con_pdf_y_tiene_aviso_y_numero_de_pagina():
    r = reporte_completo()
    datos = generar_pdf(r)
    assert datos.startswith(b"%PDF")
    texto = texto_pdf(datos)
    assert " ".join(r["aviso"].split()) in " ".join(texto.split())
    assert "Herramienta de apoyo; no sustituye el criterio clínico" in texto
    assert re.search(r"Página 1 de \d+", texto)
    assert "DIAG con" in texto and "IA" in texto
    assert "Hipótesis diagnósticas" in texto and "6A70" in texto


def test_pdf_incrusta_nunito():
    lector = PdfReader(io.BytesIO(generar_pdf(reporte_completo())))
    fuentes = {}
    for pagina in lector.pages:
        for f in (pagina["/Resources"].get("/Font") or {}).values():
            f = f.get_object()
            assert "/FontDescriptor" in f, f"fuente sin incrustar: {f['/BaseFont']}"
            fuentes[str(f["/BaseFont"])] = f["/FontDescriptor"].get_object()
    assert fuentes and all("Nunito" in nombre for nombre in fuentes)
    assert all("/FontFile2" in d for d in fuentes.values())


def test_pdf_tiene_ligas_clicables():
    lector = PdfReader(io.BytesIO(generar_pdf(reporte_completo())))
    destinos = {a.get_object()["/A"]["/URI"] for p in lector.pages for a in p.get("/Annots", [])}
    assert "https://icd.who.int/browse/2024-01/mms/es#1001" in destinos
    assert "https://www.imss.gob.mx/guias/depresion" in destinos


def test_pdf_largo_numera_varias_paginas():
    r = reporte_demo()
    r["datos_faltantes"] = [f"Dato faltante número {i} con texto de relleno suficiente para ocupar espacio" for i in range(150)]
    texto = texto_pdf(generar_pdf(r))
    assert "Página 1 de " in texto and "Página 2 de " in texto


# ----------------------------------------------------------------------------- texto hostil
def reporte_hostil() -> dict:
    hostil = "<script>alert(1)</script> &amp; ñandú áéíóú 😀🔥 \x00\x07 {{x}} <b>"
    larga = "a" * 500
    return {
        "fecha": None,
        "caso": {"paciente_ref": hostil, "edad": None, "sexo": None, "temporalidad": larga},
        "sintomas": {"cognitivos": [hostil, larga], "emocionales": [], "conductuales": None, "afectacion_social": ["😀"]},
        "hipotesis": [
            {"codigo": "6A70", "titulo": hostil, "url": "javascript:alert(1)", "justificacion": larga,
             "criterios_cumplidos": [], "criterios_no_cumplidos": [hostil, None]},
            {"codigo": None, "titulo": None, "url": None, "justificacion": None, "criterios_cumplidos": None, "criterios_no_cumplidos": []},
        ],
        "ambiguedades": [], "diferenciales": [hostil],
        "alternativas": [{"codigo": "X", "titulo": hostil, "url": "https://icd.who.int/\"><script>"}],
        "datos_faltantes": [larga, ""],
        "pruebas": [
            {"nombre": hostil, "motivo": larga, "guia": None, "cita": ""},
            {"nombre": None, "motivo": None, "guia": None, "cita": None},
            {"nombre": "p", "motivo": "m", "guia": {"titulo": hostil, "url": "https://x.org/?a=1&b=<2>"}, "cita": hostil},
        ],
        "guias": [{"tipo": "oms", "etiqueta": hostil, "titulo": larga, "url": "https://x.org", "cita": None}],
        "aviso": hostil,
    }


def reporte_vacio() -> dict:
    return {"fecha": "no es fecha", "caso": {}, "sintomas": {}, "hipotesis": [], "ambiguedades": [], "diferenciales": [],
            "alternativas": [], "datos_faltantes": [], "pruebas": [], "guias": [], "aviso": ""}


@pytest.mark.parametrize("formato", ["docx", "pdf"])
@pytest.mark.parametrize("fabrica", [reporte_hostil, reporte_vacio])
def test_exportadores_aceptan_texto_hostil(formato, fabrica):
    if formato == "docx":
        datos = generar_docx(fabrica())
        Document(io.BytesIO(datos))
    else:
        datos = generar_pdf(fabrica())
        assert len(PdfReader(io.BytesIO(datos)).pages) >= 1
    assert datos


def test_docx_conserva_el_texto_hostil_como_texto_y_sin_ligas_peligrosas():
    d = Document(io.BytesIO(generar_docx(reporte_hostil())))
    assert "<script>alert(1)</script>" in texto_docx(generar_docx(reporte_hostil()))
    ligas = [r.target_ref for r in d.part.rels.values() if r.reltype.endswith("/hyperlink")]
    assert not any(l.startswith("javascript:") for l in ligas)


def test_pdf_no_crea_ligas_peligrosas_y_muestra_el_texto_hostil():
    datos = generar_pdf(reporte_hostil())
    lector = PdfReader(io.BytesIO(datos))
    destinos = [str(a.get_object()["/A"]["/URI"]) for p in lector.pages for a in p.get("/Annots", [])]
    assert not any(d.startswith("javascript:") for d in destinos)
    assert "<script>alert(1)</script>" in texto_pdf(datos)
